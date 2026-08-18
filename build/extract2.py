# -*- coding: utf-8 -*-
"""Extractor v2 -- 當代中文課程 Book1+Book2 PPTX -> raw2.json"""
import sys, os, re, glob, json
from pptx import Presentation

# Thu muc chua 2 thu muc con cua Sach 1 / Sach 2 (cac file .pptx).
# Dat bien moi truong CCC_SOURCE, hoac sua duong dan mac dinh duoi day.
ROOT = os.environ.get("CCC_SOURCE",
                      os.path.join(os.path.dirname(os.path.abspath(__file__)), "source"))
# Tim thu muc cua tung cuon: bat ky thu muc con nao co ten chua "BOOK 1" / "BOOK 2".
def _book_dir(n):
    hits = []
    for dirpath, _dirs, files in os.walk(ROOT):
        if any(f.lower().endswith(".pptx") for f in files) and \
           ("book %d" % n) in dirpath.lower():
            hits.append(dirpath)
    return sorted(hits, key=len)[0] if hits else os.path.join(ROOT, "BOOK %d" % n)

BOOKS = {1: _book_dir(1), 2: _book_dir(2)}

TAGS = re.compile(r'[\U000E0000-\U000E01EF\uFE00-\uFE0F\u200b-\u200f\u2060\ufeff\u00ad]')
CJKC = '\u3400-\u4dbf\u4e00-\u9fff\uf900-\ufaff'
CJK  = re.compile('[' + CJKC + ']')
TONEC = ('\u0101\u00e1\u01ce\u00e0\u0113\u00e9\u011b\u00e8\u012b\u00ed\u01d0\u00ec'
         '\u014d\u00f3\u01d2\u00f2\u016b\u00fa\u01d4\u00f9\u01d6\u01d8\u01da\u01dc\u00fc'
         '\u0100\u00c1\u01cd\u00c0\u0112\u00c9\u011a\u00c8\u012a\u00cd\u01cf\u00cc'
         '\u014c\u00d3\u01d1\u00d2\u016a\u00da\u01d3\u00d9\u01d5\u01d7\u01d9\u01db\u00dc\u00ea\u0115')
TONE = re.compile('[' + TONEC + ']')
PYCH = "A-Za-z" + TONEC + r"\s'\u2019\-\u00b7./\uff0f()\u3001,"
LAT  = re.compile('^[' + PYCH + ']+$')
VNO  = re.compile(r'^(I{1,3}|IV|VI{0,3}|V|VII|VIII)\s*-\s*(\d+)\s*$')
GNO  = re.compile(r'^(I{1,3}|IV|VI{0,3}|V|VII|VIII|IX|X|XI+)\s*\.\s*(.+)$')
SPK  = re.compile('([' + CJKC + r']{1,3}[ ]?[' + CJKC + r']{0,2}[A-Za-z]{0,3})\s*[\uff1a:]\s*')
DIVIDERS = ('\u5bf9\u8bdd\u4e00','\u5bf9\u8bdd\u4e8c','\u5bf9\u8bdd\u4e09','\u77ed\u6587',
            '\u7ec3\u4e60','\u8bfe\u6587','\u5bf9\u8bdd','\u8bed\u6cd5','\u751f\u8bcd')
ROMAN = {'I':1,'II':2,'III':3,'IV':4,'V':5,'VI':6,'VII':7,'VIII':8}

def clean(s):
    s = TAGS.sub('', s or '')
    s = s.replace('\u3000', ' ').replace('\xa0', ' ').replace('\t', ' ')
    s = s.replace('\u201c', '"').replace('\u201d', '"')
    s = s.replace('\u2018', "'").replace('\u2019', "'")
    s = s.replace('\ufb01', 'fi').replace('\ufb02', 'fl')
    return re.sub(r'[ ]{2,}', ' ', s).strip()

def cjkonly(s):
    return re.sub('[^' + CJKC + ']', '', s or '')

def slide_texts(slide):
    out = []
    def walk(shapes):
        for sh in shapes:
            if sh.__class__.__name__ == 'GroupShape':
                walk(sh.shapes); continue
            try:
                if sh.has_text_frame:
                    lines = [clean(x) for x in sh.text_frame.text.split('\n')]
                    if any(lines):
                        out.append({'n': sh.name, 'tab': False, 'lines': lines,
                                    'top': sh.top or 0, 'left': sh.left or 0})
                elif getattr(sh, 'has_table', False):
                    rows = []
                    for r in sh.table.rows:
                        cells = [clean(c.text) for c in r.cells]
                        seen = []
                        for c in cells:
                            if c and c not in seen: seen.append(c)
                        if seen: rows.append(' '.join(seen))
                    if rows:
                        out.append({'n': 'TABLE', 'tab': True, 'lines': rows,
                                    'top': sh.top or 0, 'left': sh.left or 0})
            except Exception:
                pass
    walk(slide.shapes)
    out.sort(key=lambda d: (d['top'], d['left']))
    return out

# ── vocabulary ───────────────────────────────────────────────────────────────
def vocab_marker(sts):
    for s in sts:
        for l in s['lines']:
            if VNO.match(l): return VNO.match(l).group(0).replace(' ', '')
    return None

PARTICLE_PY = re.compile(r'^([a-z\u00fc]{1,6})\s*[\(\uff08]')

def parse_vocab(sts, marker):
    lines = [l for s in sts for l in s['lines'] if l]
    lines = [l for l in lines if not VNO.match(l)]
    lines = [l for l in lines if l not in DIVIDERS]
    if not lines: return None
    hanzi = None; py = None; en = []
    for l in lines:
        # a line that is CJK glued to romanisation -> split it
        m = re.match('^([' + CJKC + r'\uff08\uff09()\uff0f/\u3001\uff0b+\uff5e~\u3010\u3011]{1,14})([' + PYCH + ']{2,})$', l)
        if m and hanzi is None and CJK.search(m.group(1)) and TONE.search(m.group(2)):
            hanzi = m.group(1); py = re.sub(r'\s+', ' ', m.group(2)).strip(); continue
        cjk_n = len(CJK.findall(l))
        core = re.sub(r'[\s\u3002\uff0c\u3001\uff01\uff1f\u00b7]', '', l)
        # pure-CJK headword line (allow （） ／ ＋ ~ brackets)
        if cjk_n and hanzi is None and cjk_n <= 12 and \
           len(re.sub('[' + CJKC + r'\uff08\uff09()\uff0f/\u3001\uff0b+\uff5e~\u3010\u3011\s]', '', core)) == 0:
            hanzi = re.sub(r'\s', '', l); continue
        if py is None and not cjk_n and TONE.search(l) and LAT.match(l) and len(l) < 46:
            py = re.sub(r'\s+', ' ', l).strip(); continue
        if l: en.append(l)
    entxt = ' '.join(en).strip()
    # patterns like "\u661f\u671f\uff0b number (N) week"  -> hanzi already caught? else pull leading CJK
    if hanzi is None:
        m = re.match('^([' + CJKC + r'\uff0b+\uff08\uff09()\uff0f/]{1,14})\s*(.*)$', entxt)
        if m and CJK.search(m.group(1)):
            hanzi = re.sub(r'\s', '', m.group(1)); entxt = m.group(2).strip()
    # latin headword (KTV, 101 ...)
    if hanzi is None:
        m = re.match(r'^([A-Z0-9]{2,6})\s+(.*)$', entxt)
        if m: hanzi = m.group(1); entxt = m.group(2).strip()
    # pinyin still missing: leading romanisation of entxt (neutral-tone particles etc.)
    if py is None and entxt:
        m = PARTICLE_PY.match(entxt)
        if m:
            py = m.group(1); entxt = entxt[m.end(1):].strip()
        else:
            m = re.match('^([' + PYCH + r']{2,30}?)\s*[\(\uff08]', entxt)
            if m and not CJK.search(m.group(1)):
                py = re.sub(r'\s+', ' ', m.group(1)).strip(); entxt = entxt[m.end(1):].strip()
            else:
                m = re.match('^([' + PYCH + r']{2,30})$', entxt)
                if m and TONE.search(entxt):
                    py = re.sub(r'\s+', ' ', entxt).strip(); entxt = ''
    # POS tag
    pos = ''
    m = re.match(r'^[\(\uff08]([A-Za-z/,\.\s\-]{1,22})[\)\uff09]\s*(.*)$', entxt, re.S)
    if m:
        pos = re.sub(r'\s+', ' ', m.group(1)).strip(); entxt = m.group(2).strip()
    if py: py = py.strip(' .,')
    return {'no': marker, 'w': hanzi, 'p': py, 'pos': pos, 'en': re.sub(r'\s+', ' ', entxt).strip()}

# ── running text ─────────────────────────────────────────────────────────────
def split_speakers(block):
    """block: list of raw lines -> list of {sp, zh} or {zh} paragraphs"""
    joined = ' \u0001 '.join(block)
    hits = list(SPK.finditer(joined))
    # only treat as dialogue if markers look like real speaker turns
    if len(hits) >= 2:
        out = []
        for i, m in enumerate(hits):
            end = hits[i+1].start() if i + 1 < len(hits) else len(joined)
            body = joined[m.end():end]
            body = body.replace('\u0001', ' ')
            body = re.sub(r'\s+', ' ', body).strip()
            sp = re.sub(r'\s', '', m.group(1))
            if body: out.append({'sp': sp, 'zh': body})
        return out
    return [{'zh': re.sub(r'\s+', ' ', l).strip()} for l in block if cjkonly(l)]

def collect_text(slides, vslide, lo, hi):
    """all dialogue/prose blocks in slide range [lo,hi), deduped by containment"""
    blocks = []
    for i in range(lo, hi):
        if i in vslide: continue
        for s in slides[i]:
            ls = [l for l in s['lines'] if l and l not in DIVIDERS]
            if not ls: continue
            if GNO.match(ls[0]): continue
            cjk = sum(len(CJK.findall(l)) for l in ls)
            if cjk < 10: continue
            lat = sum(1 for l in ls if TONE.search(l))
            if lat > len(ls) / 2.0: continue                 # pinyin-annotated duplicate
            if re.match(r'^\u7b2c[\u4e00-\u9fff\d]+\u8bfe$', ls[0]): continue
            blocks.append({'slide': i, 'lines': ls, 'norm': cjkonly(' '.join(ls))})
    # dedupe by containment (keep longest)
    keep = []
    for b in sorted(blocks, key=lambda x: -len(x['norm'])):
        if any(b['norm'] in k['norm'] for k in keep): continue
        keep.append(b)
    keep.sort(key=lambda x: x['slide'])
    return keep

def extract(book, lid, path):
    prs = Presentation(path)
    slides = [slide_texts(s) for s in prs.slides]
    vocab = []; vslide = {}
    for i, sts in enumerate(slides):
        m = vocab_marker(sts)
        if not m: continue
        v = parse_vocab(sts, m)
        if v is None or (not v['w'] and not v['p']):
            vslide[i] = m; continue
        v['slide'] = i
        vocab.append(v); vslide[i] = m
    grp_first, grp_last = {}, {}
    for v in vocab:
        g = VNO.match(v['no']).group(1)
        grp_first.setdefault(g, v['slide'])
        grp_last[g] = v['slide']
    title = {'zh': None, 'py': None, 'en': None, 'no': None}
    for sts in slides[:5]:
        flat = [l for s in sts for l in s['lines'] if l]
        for j, l in enumerate(flat):
            if re.match(r'^\u7b2c[\u4e00-\u9fff]+\u8bfe$', l) or re.match(r'^\u7b2c\s*\d+\s*\u8bfe$', l):
                title['no'] = l
                for r in flat[j+1:]:
                    if CJK.search(r) and title['zh'] is None and len(r) < 26: title['zh'] = r
                    elif TONE.search(r) and title['py'] is None: title['py'] = r
                    elif LAT.match(r) and title['en'] is None and len(r) > 3: title['en'] = r
        if title['zh']: break
    order = sorted(grp_first.items(), key=lambda kv: kv[1])
    texts = []; prev_end = 0
    for g, first in order:
        blocks = collect_text(slides, vslide, prev_end, first)
        lines = []
        for b in blocks: lines.extend(b['lines'])
        turns = split_speakers(lines) if lines else []
        nsp = sum(1 for t in turns if t.get('sp'))
        texts.append({'grp': g, 'gi': ROMAN.get(g, 9), 'kind': 'dialogue' if nsp >= 2 else 'essay',
                      'nsp': nsp, 'slides': [b['slide'] for b in blocks], 'turns': turns})
        prev_end = grp_last[g] + 1
    gram = []
    for i, sts in enumerate(slides):
        if i in vslide: continue
        for s in sts:
            ls = [l for l in s['lines'] if l]
            if ls and GNO.match(ls[0]):
                gram.append({'slide': i, 'head': ls[0],
                             'body': [l for s2 in sts for l in s2['lines'] if l and l != ls[0]]})
                break
    return {'id': lid, 'book': book, 'title': title, 'vocab': vocab,
            'texts': texts, 'grammar': gram, 'nslides': len(slides)}

if __name__ == '__main__':
    out = {}
    for b, d in sorted(BOOKS.items()):
        for f in sorted(glob.glob(os.path.join(d, '*.pptx'))):
            lid = os.path.basename(f).replace('.pptx', '')
            r = extract(b, lid, f); out[lid] = r
            bad = [v for v in r['vocab'] if not v['w'] or not v['p']]
            print("%-8s %-14s vocab=%-3d bad=%-2d texts=%s gram=%d" % (
                lid, (r['title']['zh'] or '??')[:14], len(r['vocab']), len(bad),
                [(t['grp'], t['kind'][:3], len(t['turns'])) for t in r['texts']], len(r['grammar'])))
    json.dump(out, open('raw2.json', 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    print('\nTOTAL VOCAB:', sum(len(v['vocab']) for v in out.values()))
    print('TOTAL BAD  :', sum(1 for v in out.values() for x in v['vocab'] if not x['w'] or not x['p']))
    print('TOTAL TURNS:', sum(len(t['turns']) for v in out.values() for t in v['texts']))
